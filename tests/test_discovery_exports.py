import io
import httpx
from openpyxl import load_workbook
from app import config
from app.discovery import queries,normalize,canonical_url,contacts,BraveProvider
from app.criteria import Criteria,screen
from app.db import Target
from app.exports import export_targets
from app.security import safe_public_url

def test_custom_query_and_separate_category_counts():
    c=Criteria(industries=['Commercial Pool Maintenance'],mode='both',target_count=250)
    q=queries(c.model_dump());assert {'brokered','proprietary'}=={v['category'] for v in q}
    assert all('Commercial Pool Maintenance' in v['query'] for v in q)

def test_actual_brave_request_adapter_mocked(monkeypatch):
    monkeypatch.setattr(config,'BRAVE_KEY','fixture-key')
    original=httpx.Client
    def handler(r):
        assert r.url.host=='api.search.brave.com' and r.headers['X-Subscription-Token']=='fixture-key'
        return httpx.Response(200,json={'web':{'results':[{'title':'Example Services','url':'https://example.com','description':'Company services'}]}})
    monkeypatch.setattr(httpx,'Client',lambda **kw:original(transport=httpx.MockTransport(handler),**kw))
    response=BraveProvider().search('Example services');assert len(response['web']['results'])==1

def test_normalization_contacts_and_sde():
    result={'title':'Maintenance business for sale','url':'https://broker.example/listing/123','description':'Revenue: $5M EBITDA: $900K SDE: $1M Asking price: $4M. Contact broker@example.com or 604-555-0101'}
    d=normalize(result,'brokered','maintenance');assert d['company_name'] is None
    assert d['facts']['ebitda']['value']==900000 and d['facts']['sde']['value']==1000000
    assert len(d['contacts'])==2 and d['contacts'][0]['verification'].endswith('unverified')
    assert screen(d,Criteria())['status']=='Needs Verification'

def test_dedup_keeps_listings_separate():
    assert canonical_url('https://www.example.com/about?utm_source=x','proprietary')==canonical_url('https://example.com','proprietary')
    assert canonical_url('https://example.com/listing/1','brokered')!=canonical_url('https://example.com/listing/2','brokered')
    assert canonical_url('https://www.example.com/listing/1?utm_source=x','brokered')==canonical_url('https://example.com/listing/1','brokered')

def test_no_unsafe_urls_or_fake_contacts():
    for url in ['http://example.com','https://127.0.0.1','https://169.254.169.254','https://a.internal','https://user:pass@example.com']:
        assert safe_public_url(url) is None
    assert contacts('Owner name unavailable.','https://example.com')==[]

def test_excel_and_csv_no_sourcing_scores_and_formula_safety():
    target=Target(category='brokered',data={'name':'=HYPERLINK("bad")','url':'https://example.com','export_allowed':True,'facts':{},'contacts':[]},screening={'status':'Needs Verification','missing':['Country']})
    raw=export_targets([target],'xlsx');wb=load_workbook(io.BytesIO(raw));assert wb.active['A2'].value.startswith("'")
    assert not any('score' in str(cell.value).lower() for cell in wb.active[1])
    csv=export_targets([target],'csv').decode('utf-8-sig');assert 'Needs Verification' in csv
    target.data={**target.data,'export_allowed':False};assert wb.active.max_row==2
    assert load_workbook(io.BytesIO(export_targets([target],'xlsx'))).active.max_row==1


def test_explicit_contact_names_and_linkedin_only():
    c=contacts('CEO: John Smith. https://www.linkedin.com/in/john-smith','https://example.com')
    assert c[0]['name']=='John Smith' and c[0]['title']=='CEO'
    assert c[1]['linkedin']=='https://www.linkedin.com/in/john-smith'
