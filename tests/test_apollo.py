import copy
import io
import json
from pathlib import Path
import httpx
import pytest
from sqlalchemy import select,func
from openpyxl import load_workbook
from app import config
from app.apollo import ApolloProvider,build_org_filters,revenue_band,normalize_org,normalize_people,independent,INDUSTRY_MAP
from app.criteria import Criteria,screen
from app.db import Session,Job,Target,Usage,ProviderRate,User
from app.jobs import enqueue,tick
from app.pipeline import provider_rate
from app.providers import ProviderThrottled,ProviderUnavailable,TransientProviderError
from app.exports import export_targets

FIXTURES=Path(__file__).parent/'fixtures'
def fixture(name):return json.loads((FIXTURES/f'apollo_{name}.json').read_text())

@pytest.fixture
def apollo_config(monkeypatch):
    monkeypatch.setattr(config,'APOLLO_KEY','test-fixture-key')
    monkeypatch.setattr(config,'APOLLO_CREDIT_COST',.01)
    monkeypatch.setattr(config,'APOLLO_STORAGE_RIGHTS',True)
    monkeypatch.setattr(config,'APOLLO_EXPORT_RIGHTS',True)
    monkeypatch.setattr(config,'APOLLO_SEARCH_CREDITS',1)
    monkeypatch.setattr(config,'APOLLO_ENRICH_CREDITS',1)
    monkeypatch.setattr(config,'APOLLO_REVEAL_CREDITS',1)

@pytest.fixture
def provider_mocks(apollo_config,monkeypatch):
    calls=[]
    def mock_search(self,filters,page=1):calls.append(('search',page));return fixture('search')
    def mock_enrich(self,domain):calls.append(('enrich',domain));return fixture('enrich')
    def mock_people(self,org_id):calls.append(('people',org_id));return fixture('people')
    def mock_reveal(self,person_id):calls.append(('reveal',person_id));return fixture('reveal')
    monkeypatch.setattr(ApolloProvider,'search',mock_search);monkeypatch.setattr(ApolloProvider,'enrich',mock_enrich);monkeypatch.setattr(ApolloProvider,'people',mock_people);monkeypatch.setattr(ApolloProvider,'reveal',mock_reveal)
    monkeypatch.setattr('app.pipeline.provider_rate',lambda db:None)
    return calls

def test_all_questionnaire_industries_mapped():
    import re
    js=(config.ROOT/'app/static/app.js').read_text()
    options=re.search(r'const INDUSTRIES=\[(.*?)\];',js).group(1)
    assert set(re.findall(r"'([^']+)'",options))==set(INDUSTRY_MAP)

def test_criteria_mapping_and_custom_industries():
    c=Criteria(industries=['HVAC'],exclusions=['Restaurants'],countries=['United States'],regions=['Ohio'],employees={'minimum':10,'maximum':100},revenue={'minimum':3000000,'maximum':20000000})
    p=build_org_filters(c)
    assert p['organization_locations[]']==['ohio, united states']
    assert p['organization_num_employees_ranges[]']==['10,100']
    assert p['revenue_range[min]']==3000000
    assert p['organization_naics_codes[]']==['238220']
    assert 'hvac' in p['q_organization_keyword_tags[]']
    assert p['not_organization_naics_codes[]']==['7225']
    custom=build_org_filters(Criteria(industries=['Specialty undersea sensors']))
    assert custom['q_organization_keyword_tags[]']==['Specialty undersea sensors']
    assert 'organization_naics_codes[]' not in custom
    assert build_org_filters(Criteria(employees={'minimum':0,'maximum':10}))['organization_num_employees_ranges[]']==['0,10']

def test_shared_naics_exclusion_does_not_erase_selected_industry():
    p=build_org_filters(Criteria(industries=['HVAC'],exclusions=['Plumbing']))
    assert 'not_organization_naics_codes[]' not in p
    assert p['q_not_organization_keyword_tags[]']==['plumbing']

def test_ebitda_only_widens_search_and_cad_not_used_as_usd():
    c=Criteria(revenue={'minimum':5e6,'maximum':10e6},ebitda={'minimum':1e6,'maximum':2e6})
    lo,hi,assumption=revenue_band(c)
    assert lo<=5e6 and hi>=10e6 and assumption['basis'].startswith('Illustrative')
    assert c.revenue.maximum==10e6
    p=build_org_filters(Criteria(currency='CAD',revenue={'minimum':1e6}))
    assert 'revenue_range[min]' not in p

def test_pre_filters_and_facts_are_estimates(apollo_config):
    rows=fixture('search')['organizations']
    assert independent(rows[0])[0]
    assert not independent(rows[2])[0] and not independent(rows[3])[0]
    d=normalize_org(fixture('enrich')['organization'])
    assert d['facts']['revenue']['status']=='reported'
    assert d['facts']['revenue']['basis']=='modeled_estimate'
    assert d['facts']['employees']['basis']=='modeled_estimate'
    assert d['facts']['country']['provider']=='Apollo'
    assert d['apollo_id']=='fixture-independent'
    assert d['facts']['primary_phone']['value']=='+1 614-555-0101'

def test_screening_tiers_and_no_numeric_score(apollo_config):
    d=normalize_org(fixture('enrich')['organization']);c=Criteria(industries=['HVAC'],countries=['United States'],regions=['Ohio'])
    r=screen(d,c);assert r['tier']=='Strong fit' and 'fit_score' not in r and r['reported']
    c.revenue.minimum=10e6;r=screen(d,c);assert r['tier']=='Likely out of range' and r['status']!='Excluded'
    c.revenue.minimum=1e6;r=screen(d,c);assert r['tier']=='Possible fit'
    d['facts']['country']['value']='Canada';assert screen(d,c)['tier']=='Excluded'

def test_funding_preference_and_parent_exclusion(apollo_config):
    org=fixture('enrich')['organization'];org['total_funding']=1000000;org['latest_funding_stage']='series_a'
    d=normalize_org(org);assert 'Institutionally funded' in d['flags']
    assert screen(d,Criteria())['status']!='Excluded'
    assert screen(d,Criteria(characteristics={'private_ownership':'Required'}))['status']=='Excluded'
    org['ultimate_parent_organization']={'name':'Parent'}
    assert screen(normalize_org(org),Criteria())['status']=='Excluded'

def test_people_does_not_reveal_contacts():
    rows=normalize_people(fixture('people'),'https://app.apollo.io')
    assert rows[0]['name']=='Alex' and 'email' not in rows[0] and not rows[0]['revealed']
    assert rows[0]['has_email'] and rows[0]['verification'].endswith('not revealed')

def test_http_adapters_shape_and_errors(apollo_config,monkeypatch):
    original=httpx.Client;requests=[]
    def handler(r):
        requests.append(r)
        assert r.url.host=='api.apollo.io' and r.headers['x-api-key']=='test-fixture-key'
        kind='search' if r.url.path.endswith('mixed_companies/search') else 'enrich' if r.url.path.endswith('organizations/enrich') else 'people' if r.url.path.endswith('mixed_people/api_search') else 'reveal'
        return httpx.Response(200,json=fixture(kind))
    monkeypatch.setattr(httpx,'Client',lambda **kw:original(transport=httpx.MockTransport(handler),**kw))
    provider=ApolloProvider();provider.search(build_org_filters(Criteria(industries=['HVAC'])),page=2);provider.enrich('hvac.example.com');provider.people('fixture-independent');provider.reveal('fixture-owner')
    assert requests[0].url.params['page']=='2' and requests[0].url.params['per_page']=='100'
    assert requests[0].url.params.get_list('organization_naics_codes[]')==['238220']
    assert json.loads(requests[2].content)['person_seniorities']==['owner','founder','c_suite']
    assert json.loads(requests[3].content)['reveal_phone_number'] is False

@pytest.mark.parametrize('status,error',[(403,ProviderUnavailable),(429,TransientProviderError),(500,TransientProviderError)])
def test_plan_and_transient_errors(apollo_config,monkeypatch,status,error):
    original=httpx.Client
    monkeypatch.setattr(httpx,'Client',lambda **kw:original(transport=httpx.MockTransport(lambda _:httpx.Response(status)),**kw))
    with pytest.raises(error):ApolloProvider().search({})

def test_shared_rate_limit_enforced(apollo_config):
    with Session() as db:
        provider_rate(db);db.commit()
        with pytest.raises(ProviderThrottled):provider_rate(db)
        record=db.scalar(select(ProviderRate));assert record.calls==1
        record.next_allowed=0;record.calls=config.APOLLO_CALLS_PER_HOUR;db.commit()
        with pytest.raises(ProviderThrottled):provider_rate(db)

def test_pipeline_enriches_one_deduped_company_no_investment_evaluation(client,provider_mocks):
    body={'criteria':Criteria(mode='proprietary',industries=['HVAC'],countries=['United States'],regions=['Ohio']).model_dump(),'budget':1,'idempotency_key':'apollo-flow'}
    r=client.post('/api/search',json=body);assert r.status_code==200
    assert tick()
    targets=client.get('/api/targets?category=proprietary').json();assert len(targets)==1
    assert targets[0]['screening']['tier']=='Strong fit'
    assert targets[0]['data']['_apollo_stage']=='complete'
    assert provider_mocks==[('search',1),('enrich','hvac.example.com'),('people','fixture-independent')]
    assert targets[0]['data']['contacts'][-1]['name']=='Alex'
    with Session() as db:
        usage=db.scalars(select(Usage).order_by(Usage.id)).all()
        assert len(usage)==3 and sum(u.estimated_cost for u in usage)==.02
        assert all(u.detail['response_received'] for u in usage)
    assert all(j['kind']=='sourcing' for j in client.get('/api/jobs').json())

def test_reveal_is_manual_reserved_and_tenant_scoped(client,provider_mocks):
    client.post('/api/search',json={'criteria':Criteria(mode='proprietary').model_dump(),'budget':1,'idempotency_key':'source-before-reveal'});tick()
    target=client.get('/api/targets?category=proprietary').json()[0];tid=target['id']
    assert not any(k=='reveal' for k,_ in provider_mocks)
    r=client.post(f'/api/targets/{tid}/reveal',json={'person_id':'fixture-owner','budget':.02,'idempotency_key':'explicit-reveal'})
    assert r.status_code==200 and r.json()['kind']=='contact_reveal';tick()
    profile=client.get(f'/api/targets/{tid}').json()
    assert profile['data']['contacts'][-1]['email']=='alex@example.com'
    assert profile['data']['contacts'][-1]['revealed']
    assert provider_mocks[-1]==('reveal','fixture-owner')
    with Session() as db:
        u=db.scalar(select(Usage).where(Usage.category=='contact_enrichment'));assert u.estimated_cost==.01
    assert client.post(f'/api/targets/{tid}/reveal',json={'person_id':'someone-else','budget':.02,'idempotency_key':'foreign-contact'}).status_code==404

def test_budget_stops_before_enrichment_and_retry_reuses_search(client,provider_mocks):
    r=client.post('/api/search',json={'criteria':Criteria(mode='proprietary').model_dump(),'budget':.015,'idempotency_key':'low-apollo-budget'})
    tick();job=client.get('/api/jobs').json()[0]
    assert job['status']=='failed' and job['spent']==.01
    assert provider_mocks==[('search',1)]
    with Session() as db:
        j=db.get(Job,r.json()['id']);j.budget=1;j.status='queued';db.commit()
    tick();assert provider_mocks.count(('search',1))==1
    assert len(client.get('/api/targets?category=proprietary').json())==1

def test_estimate_separates_sources_and_discloses_defaults(client,apollo_config,monkeypatch):
    monkeypatch.setattr(config,'BRAVE_KEY','fixture');monkeypatch.setattr(config,'BRAVE_STORAGE_RIGHTS',True)
    r=client.post('/api/search/estimate',json=Criteria(mode='both',target_count=250).model_dump()).json()
    assert r['apollo_enrichments']==250 and r['apollo_search_pages']==9 and r['brave_queries']>0
    assert r['apollo_credits']==259 and r['available']
    boot=client.get('/api/bootstrap').json();assert boot['defaults']=={'search_budget':10,'monthly_budget':100}

def test_export_labels_estimates_and_omits_fit_scores(apollo_config):
    d=normalize_org(fixture('enrich')['organization']);d['contacts']=normalize_people(fixture('people'),d['apollo_url'])
    t=Target(category='proprietary',data=d,screening=screen(d,Criteria()))
    wb=load_workbook(io.BytesIO(export_targets([t],'xlsx')));headers=[c.value for c in wb.active[1]]
    assert 'Fit tier' in headers and 'Fit reasons' in headers and 'Employees (Apollo est.)' in headers
    assert not any('score' in h.lower() for h in headers)
    d['export_allowed']=False;assert load_workbook(io.BytesIO(export_targets([t],'xlsx'))).active.max_row==1


def test_throttled_job_does_not_count_as_provider_failure(client,apollo_config,monkeypatch):
    monkeypatch.setattr('app.pipeline.provider_rate',lambda _:(_ for _ in ()).throw(ProviderThrottled('Paced request',6)))
    r=client.post('/api/search',json={'criteria':Criteria(mode='proprietary').model_dump(),'budget':1,'idempotency_key':'throttle-is-not-failure'})
    tick()
    with Session() as db:
        j=db.get(Job,r.json()['id']);assert j.status=='retrying' and j.attempts==0 and j.spent==0
        assert db.scalar(select(func.count()).select_from(Usage))==0


def test_both_categories_interleaved(client,provider_mocks,monkeypatch):
    monkeypatch.setattr(config,'BRAVE_KEY','fixture')
    monkeypatch.setattr('app.pipeline.queries',lambda _:[{'category':'brokered','query':'sale','offset':0}])
    monkeypatch.setattr('app.pipeline.BraveProvider.search',lambda *args,**kwargs:{'web':{'results':[{'title':'Fixture business for sale','url':'https://broker.example.com/one','description':'Revenue: $5M'}]}})
    r=client.post('/api/search',json={'criteria':Criteria(mode='both').model_dump(),'budget':1,'idempotency_key':'both-providers'})
    assert r.status_code==200;tick()
    jobs=client.get('/api/jobs').json();assert jobs[0]['progress']=={'brokered':1,'proprietary':1}
    assert len(client.get('/api/targets?category=brokered').json())==1
    assert len(client.get('/api/targets?category=proprietary').json())==1


def test_cached_provider_result_does_not_reserve_again(user,provider_mocks):
    from app.pipeline import apollo_call
    with Session() as db:
        job=enqueue(db,user,'sourcing',{'criteria':Criteria(mode='proprietary').model_dump()},'cache-credit-test',1)
        first=apollo_call(db,job,'search',{'filters':{},'page':1},1)
        second=apollo_call(db,job,'search',{'filters':{},'page':1},1)
        assert first==second and job.spent==.01
        assert db.scalar(select(func.count()).select_from(Usage))==1
    assert provider_mocks==[('search',1)]


def test_changed_domain_from_enrichment_not_attached(client,provider_mocks,monkeypatch):
    wrong=fixture('enrich');wrong['organization']['website_url']='https://wrong.example.com';wrong['organization']['primary_domain']='wrong.example.com'
    monkeypatch.setattr(ApolloProvider,'enrich',lambda self,domain:wrong)
    client.post('/api/search',json={'criteria':Criteria(mode='proprietary').model_dump(),'budget':1,'idempotency_key':'wrong-company-rejection'});tick()
    job=client.get('/api/jobs').json()[0];assert job['status']=='failed' and 'different company domain' in job['error']
    target=client.get('/api/targets?category=proprietary').json()[0]
    assert target['data']['primary_domain']=='hvac.example.com' and target['data']['_apollo_stage']=='pending'


def test_reveal_cannot_access_another_users_target(client,user,provider_mocks):
    with Session() as db:
        other=User(email='apollo-other@example.com');db.add(other);db.flush()
        job=Job(user_id=other.id,kind='sourcing',status='completed',payload={'criteria':Criteria().model_dump()},idempotency_key='other-job');db.add(job);db.flush()
        d=normalize_org(fixture('enrich')['organization']);d['contacts']=normalize_people(fixture('people'),d['apollo_url'])
        target=Target(user_id=other.id,job_id=job.id,category='proprietary',identity='hvac.example.com',data=d,screening=screen(d,Criteria()));db.add(target);db.commit();tid=target.id
    assert client.post(f'/api/targets/{tid}/reveal',json={'person_id':'fixture-owner','budget':1,'idempotency_key':'not-my-contact'}).status_code==404
    assert not provider_mocks


def test_cad_filters_do_not_exclude_usd_estimate_without_conversion(apollo_config):
    d=normalize_org(fixture('enrich')['organization'])
    r=screen(d,Criteria(currency='CAD',revenue={'minimum':1e6,'maximum':2e6}))
    assert r['tier']=='Possible fit' and r['status']!='Excluded'
    assert any('currency' in m.lower() for m in r['missing'])


def test_funding_amount_alone_does_not_prove_institutional_ownership(apollo_config):
    org=fixture('enrich')['organization'];org['total_funding']=100000;org['latest_funding_stage']='grant'
    d=normalize_org(org)
    assert 'institutionally_funded' not in d['facts']
    assert 'Funding reported; investor type needs verification' in d['flags']
    assert not any('Institutional funding conflicts' in reason for reason in screen(d,Criteria(characteristics={'private_ownership':'Required'}))['conflicts'])


def test_tier_filters_apply_to_dashboard_and_spreadsheet(client,provider_mocks):
    client.post('/api/search',json={'criteria':Criteria(mode='proprietary').model_dump(),'budget':1,'idempotency_key':'tier-export-source'});tick()
    assert len(client.get('/api/targets?category=proprietary&tier=Strong%20fit').json())==1
    assert client.get('/api/targets?category=proprietary&tier=Excluded').json()==[]
    csv=client.get('/api/export/proprietary/csv?tier=Strong%20fit').text
    assert 'Example HVAC Services' in csv and 'Fit tier' in csv
    assert 'Example HVAC Services' not in client.get('/api/export/proprietary/csv?tier=Excluded').text


def test_legacy_job_does_not_switch_provider_when_apollo_enabled(user,apollo_config,monkeypatch):
    monkeypatch.setattr(config,'BRAVE_KEY','fixture')
    monkeypatch.setattr('app.jobs.queries',lambda _:[{'category':'proprietary','query':'legacy','offset':0}])
    calls=[]
    monkeypatch.setattr('app.jobs.BraveProvider.search',lambda *args,**kwargs:calls.append('brave') or {'web':{'results':[]}})
    monkeypatch.setattr(ApolloProvider,'search',lambda *args,**kwargs:(_ for _ in ()).throw(AssertionError('An existing job must not change providers.')))
    with Session() as db:enqueue(db,user,'sourcing',{'criteria':Criteria(mode='proprietary').model_dump()},'legacy-provider-snapshot',1)
    tick();assert calls==['brave']
