import csv
import hashlib
import io
import zipfile
from pathlib import Path
from openpyxl import load_workbook
from pypdf import PdfReader
from . import config
from .security import cipher

ALLOWED = {'.pdf', '.xlsx', '.csv', '.txt'}

def validate(filename, content):
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED: raise ValueError('Supported files: PDF, XLSX, CSV, TXT. Legacy XLS and macros are not accepted.')
    if not content or len(content) > config.MAX_UPLOAD: raise ValueError('Files must be nonempty and at most 10 MB.')
    if suffix == '.pdf' and not content.startswith(b'%PDF-'): raise ValueError('Invalid PDF signature.')
    if suffix == '.xlsx':
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if '[Content_Types].xml' not in archive.namelist() or sum(i.file_size for i in archive.infolist()) > 50*1024*1024: raise ValueError('Invalid or oversized spreadsheet.')
        except zipfile.BadZipFile: raise ValueError('Invalid spreadsheet.')
    if suffix in {'.csv','.txt'}:
        try: content.decode('utf-8-sig')
        except UnicodeDecodeError: raise ValueError('Text files must be UTF-8.')
    return suffix


def store(filename, content):
    validate(filename, content)
    import secrets
    key = secrets.token_hex(24)
    (config.DATA / 'documents' / key).write_bytes(cipher().encrypt(content))
    (config.DATA / 'documents' / key).chmod(0o600)
    return key, hashlib.sha256(content).hexdigest()


def extract(filename, content):
    suffix = validate(filename, content)
    chunks = []
    if suffix == '.pdf':
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted: raise ValueError('Password-protected PDFs must be unlocked before upload.')
        if len(reader.pages) > 300: raise ValueError('PDF exceeds 300-page limit.')
        for i,page in enumerate(reader.pages):
            text = page.extract_text() or ''
            if text.strip(): chunks.append({'location': f'page {i+1}', 'text': text})
        if not chunks: raise ValueError('No extractable text. Upload a text PDF or an OCR-processed copy.')
    elif suffix == '.xlsx':
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
        count = 0
        for sheet in wb:
            for i,row in enumerate(sheet.iter_rows(values_only=True),1):
                count += 1
                if count > 10000: raise ValueError('Spreadsheet exceeds 10,000-row extraction limit.')
                if any(v is not None for v in row): chunks.append({'location': f'{sheet.title}!row {i}', 'text': ' | '.join(str(v) if v is not None else '' for v in row)})
        wb.close()
    elif suffix == '.csv':
        for i,row in enumerate(csv.reader(io.StringIO(content.decode('utf-8-sig'))),1):
            if i > 10000: raise ValueError('CSV exceeds 10,000-row extraction limit.')
            chunks.append({'location': f'row {i}', 'text': ' | '.join(row)})
    else:
        for i,line in enumerate(content.decode('utf-8-sig').splitlines(),1):
            if line.strip(): chunks.append({'location': f'line {i}', 'text': line})
    if sum(len(c['text']) for c in chunks) > 500000: raise ValueError('Document exceeds text extraction limit.')
    return chunks


def read_document(document):
    path = config.DATA / 'documents' / document.storage
    return cipher().decrypt(path.read_bytes())
