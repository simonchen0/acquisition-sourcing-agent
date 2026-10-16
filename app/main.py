import csv
import io
import json
import math
import secrets
import threading
from contextlib import asynccontextmanager
from pathlib import Path
from fastapi import FastAPI, Request, HTTPException, Depends, UploadFile, File, Form
from fastapi.responses import HTMLResponse, Response, JSONResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from pydantic import BaseModel, Field, ValidationError
from sqlalchemy import select, func, delete
from sqlalchemy.orm import Session as DBSession
from . import config
from .db import init_db, local_user, Session, User, CriteriaRecord, Job, Target, Opportunity, Document, Evaluation, Usage
from .security import safe_public_url
from .criteria import Criteria, CHARACTERISTICS, screen
from .jobs import enqueue, ACTIVE, month, work
from .discovery import queries, canonical_url
from .documents import store, read_document, extract
from .evaluation import FRAMEWORKS, FinancialDatum, weights_valid
from .exports import export_targets, evaluation_markdown

@asynccontextmanager
async def lifespan(app):
    init_db()
    stop=threading.Event(); worker=None
    if config.EMBEDDED_WORKER:
        worker=threading.Thread(target=work,args=(stop,),name='job-worker',daemon=True); worker.start()
    yield
    stop.set()
    if worker: worker.join(timeout=10)

app=FastAPI(title='Acquisition Sourcing Agent',lifespan=lifespan,docs_url=None,redoc_url=None)
app.mount('/static',StaticFiles(directory=config.ROOT/'app/static'),name='static')
templates=Jinja2Templates(directory=config.ROOT/'app/templates')
# No sign-in: the app binds to loopback. A per-process token on state-changing calls plus a
# Host check stop other websites (or DNS rebinding) from driving paid API calls.
CSRF=secrets.token_urlsafe(32)
LOCAL_HOSTS={'127.0.0.1','localhost','[::1]','testserver'}

def session_db():
    with Session() as db: yield db

def owned(db,model,record_id,user_id):
    obj=db.scalar(select(model).where(model.id==record_id,model.user_id==user_id))
    if not obj: raise HTTPException(404,'Record not found.')
    return obj

def uid(request): return request.state.user.id

def public_job(job):
    return {'id':job.id,'kind':job.kind,'status':job.status,'progress':job.progress,'cursor':job.cursor,'error':job.error,'spent':round(job.spent,4),'budget':job.budget,'cancel_requested':job.cancel,'created':job.created,'criteria':job.payload.get('criteria'),'opportunity_id':job.payload.get('opportunity_id')}

@app.middleware('http')
async def controls(request,call_next):
    length=request.headers.get('content-length')
    if length and (not length.isdigit() or int(length)>11*1024*1024): return JSONResponse({'detail':'Request exceeds 11 MB limit; upload documents individually.'},413)
    if (request.headers.get('host','').rsplit(':',1)[0] or '').lower() not in LOCAL_HOSTS: return JSONResponse({'detail':'Local access only.'},403)
    if request.method not in ['GET','HEAD','OPTIONS'] and not secrets.compare_digest(request.headers.get('X-CSRF-Token',''),CSRF): return JSONResponse({'detail':'Invalid request token. Reload the page.'},403)
    if not (request.url.path in ['/health','/favicon.ico'] or request.url.path.startswith('/static/')):
        with Session() as db: request.state.user=local_user(db)
    response=await call_next(request)
    response.headers['X-Content-Type-Options']='nosniff'
    response.headers['X-Frame-Options']='DENY'
    response.headers['Referrer-Policy']='no-referrer'
    response.headers['Content-Security-Policy']="default-src 'self'; script-src 'self'; style-src 'self'; img-src 'self' data:; object-src 'none'; base-uri 'none'; frame-ancestors 'none'; form-action 'self'"
    if not request.url.path.startswith('/static/'): response.headers['Cache-Control']='no-store'
    return response

@app.get('/favicon.ico')
def favicon(): return Response(status_code=204)

@app.get('/health')
def health(db:DBSession=Depends(session_db)):
    db.execute(select(1)); return {'status':'ok','database':'connected'}

@app.get('/',response_class=HTMLResponse)
@app.get('/find',response_class=HTMLResponse)
@app.get('/deals/{category}',response_class=HTMLResponse)
@app.get('/business/{record_id}',response_class=HTMLResponse)
@app.get('/evaluate',response_class=HTMLResponse)
@app.get('/evaluate/{record_id}',response_class=HTMLResponse)
@app.get('/saved',response_class=HTMLResponse)
@app.get('/settings',response_class=HTMLResponse)
def page(request:Request):
    return templates.TemplateResponse(request=request,name='app.html',context={'csrf':CSRF})

@app.get('/api/bootstrap')
def bootstrap(request:Request,db:DBSession=Depends(session_db)):
    return {'characteristics':CHARACTERISTICS,'frameworks':FRAMEWORKS,'criteria':[{'id':c.id,'name':c.name,'data':c.data} for c in db.scalars(select(CriteriaRecord).where(CriteriaRecord.user_id==uid(request)))],
            'settings':{**config.DEFAULT_SETTINGS,**request.state.user.settings},'providers':{'brave':bool(config.BRAVE_KEY) and config.BRAVE_STORAGE_RIGHTS,'openai':bool(config.AI_KEY)},
            'pricing':{'brave_request':config.BRAVE_REQUEST_COST,'ai_input_million':config.AI_INPUT_COST,'ai_output_million':config.AI_OUTPUT_COST,'model':config.AI_MODEL},
            'month_spend':request.state.user.spend if request.state.user.month==month() else 0}

@app.post('/api/criteria')
def save_criteria(c:Criteria,request:Request,db:DBSession=Depends(session_db)):
    obj=CriteriaRecord(user_id=uid(request),name=c.name,data=c.model_dump()); db.add(obj); db.commit(); return {'id':obj.id}

@app.put('/api/criteria/{cid}')
def edit_criteria(cid:int,c:Criteria,request:Request,db:DBSession=Depends(session_db)):
    obj=owned(db,CriteriaRecord,cid,uid(request)); obj.name=c.name; obj.data=c.model_dump(); db.commit(); return {'id':obj.id}

class SearchInput(BaseModel):
    criteria: Criteria
    idempotency_key: str=Field(min_length=8,max_length=64)
    budget:float=Field(gt=0,le=100,allow_inf_nan=False)

@app.post('/api/search/estimate')
def search_estimate(c:Criteria,request:Request):
    steps=queries(c.model_dump()); amount=len(steps)*config.BRAVE_REQUEST_COST
    return {'queries':len(steps),'estimated_max_cost':round(amount,3),'categories':2 if c.mode=='both' else 1,'available':bool(config.BRAVE_KEY) and config.BRAVE_STORAGE_RIGHTS}

@app.post('/api/search')
def start_search(body:SearchInput,request:Request,db:DBSession=Depends(session_db)):
    if not config.BRAVE_KEY or not config.BRAVE_STORAGE_RIGHTS: raise HTTPException(409,'Automatic discovery needs BRAVE_SEARCH_API_KEY and confirmation that your plan permits storage (BRAVE_STORAGE_RIGHTS_CONFIRMED=true). No paid request was made.')
    limit=float(request.state.user.settings.get('search_budget',config.DEFAULT_SETTINGS['search_budget']))
    if body.budget>limit: raise HTTPException(400,'Budget exceeds your per-search limit. Adjust Settings first.')
    job=enqueue(db,request.state.user,'sourcing',{'criteria':body.criteria.model_dump()},body.idempotency_key,body.budget)
    return public_job(job)

@app.get('/api/jobs')
def job_list(request:Request,db:DBSession=Depends(session_db)):
    return [public_job(j) for j in db.scalars(select(Job).where(Job.user_id==uid(request)).order_by(Job.id.desc()).limit(100))]

@app.post('/api/jobs/{jid}/cancel')
def cancel(jid:int,request:Request,db:DBSession=Depends(session_db)):
    j=owned(db,Job,jid,uid(request))
    if j.status in ACTIVE: j.cancel=True
    db.commit(); return public_job(j)

@app.post('/api/jobs/{jid}/retry')
def retry(jid:int,request:Request,db:DBSession=Depends(session_db)):
    j=owned(db,Job,jid,uid(request))
    if j.status not in ['failed','cancelled','incomplete']: raise HTTPException(409,'Only a stopped job can be retried.')
    if j.spent>=j.budget: raise HTTPException(409,'Original approved budget exhausted. Start a new search with a revised budget.')
    j.status='queued'; j.cancel=False; j.error=None; j.attempts=0; j.lease=0; db.commit(); return public_job(j)

def visible_targets(db,user_id,category,job_id=None,include_verification=True,excluded=False):
    q=select(Target).where(Target.user_id==user_id,Target.category==category)
    if job_id: q=q.where(Target.job_id==job_id)
    results=db.scalars(q.order_by(Target.id.desc())).all()
    statuses=['Excluded'] if excluded else ['Matches Known Criteria','Needs Verification'] if include_verification else ['Matches Known Criteria']
    return [t for t in results if t.screening['status'] in statuses]

@app.get('/api/targets')
def targets(request:Request,category:str='brokered',job_id:int|None=None,verification:bool=True,excluded:bool=False,sort:str='name',db:DBSession=Depends(session_db)):
    if category not in ['brokered','proprietary']: raise HTTPException(400,'Invalid category.')
    rows=visible_targets(db,uid(request),category,job_id,verification,excluded)
    if sort in ['revenue','ebitda','employees']:
        rows.sort(key=lambda t:(t.data.get('facts',{}).get(sort,{}).get('value') is None,-float(t.data.get('facts',{}).get(sort,{}).get('value',0))))
    elif sort=='completeness': rows.sort(key=lambda t:len(t.data.get('facts',{})),reverse=True)
    elif sort in ['industry','country','location']: rows.sort(key=lambda t:str(t.data.get('facts',{}).get(sort,{}).get('value','')).lower())
    else: rows.sort(key=lambda t:str(t.data.get(sort,'')).lower())
    return [{'id':t.id,'job_id':t.job_id,'category':t.category,'data':t.data,'screening':t.screening} for t in rows]

@app.get('/api/targets/{tid}')
def target_profile(tid:int,request:Request,db:DBSession=Depends(session_db)):
    t=owned(db,Target,tid,uid(request))
    opportunities=db.scalars(select(Opportunity).where(Opportunity.user_id==uid(request),Opportunity.target_id==tid)).all()
    return {'id':t.id,'category':t.category,'data':t.data,'screening':t.screening,'opportunities':[{'id':o.id,'name':o.name} for o in opportunities], 'investment_status':'Not Evaluated' if not any(db.scalar(select(Evaluation.id).where(Evaluation.opportunity_id==o.id)) for o in opportunities) else 'See explicit evaluation versions'}

@app.get('/api/export/{category}/{kind}')
def target_export(category:str,kind:str,request:Request,job_id:int|None=None,verification:bool=True,db:DBSession=Depends(session_db)):
    if category not in ['brokered','proprietary'] or kind not in ['xlsx','csv']: raise HTTPException(400,'Invalid export.')
    rows=visible_targets(db,uid(request),category,job_id,verification)
    return Response(export_targets(rows,kind),media_type='text/csv' if kind=='csv' else 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="{category}-targets.{kind}"'})

@app.post('/api/import')
def import_targets(request:Request,category:str=Form(...),criteria_id:int=Form(...),file:UploadFile=File(...),db:DBSession=Depends(session_db)):
    if category not in ['brokered','proprietary']: raise HTTPException(400,'Invalid category.')
    criteria=owned(db,CriteriaRecord,criteria_id,uid(request))
    raw=file.file.read(config.MAX_UPLOAD+1)
    if len(raw)>config.MAX_UPLOAD: raise HTTPException(413,'Import exceeds 10 MB.')
    try:
        if Path(file.filename or '').suffix.lower()=='.xlsx':
            from openpyxl import load_workbook
            from .documents import validate
            validate(file.filename,raw); wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True,keep_links=False); it=wb.active.iter_rows(values_only=True); headers=[str(v or '').strip().lower() for v in next(it)]; rows=[dict(zip(headers,row)) for row in __import__('itertools').islice(it,2001)]; wb.close()
        elif Path(file.filename or '').suffix.lower()=='.csv': rows=list(__import__('itertools').islice(csv.DictReader(io.StringIO(raw.decode('utf-8-sig'))),2001))
        else: raise ValueError('Import CSV or XLSX.')
        if len(rows)>2000: raise ValueError('Imports are limited to 2,000 rows.')
        job=Job(user_id=uid(request),kind='sourcing',status='completed',payload={'criteria':criteria.data,'import':True},idempotency_key=secrets.token_hex(16))
        db.add(job); db.flush(); count=0; seen=set()
        for row in rows:
            name=str(row.get('name') or '').strip(); url=safe_public_url(str(row.get('url') or row.get('website') or ''))
            if not name or not url: continue
            identity=canonical_url(url,category)
            if identity in seen: continue
            seen.add(identity)
            facts={}
            for key in ['industry','country','location','currency','revenue','ebitda','sde','asking_price','employees']:
                v=row.get(key)
                if v is None or v=='': continue
                if key in ['revenue','ebitda','sde','asking_price','employees']:
                    v=float(str(v).replace(',','').replace('$',''))
                    if not math.isfinite(v) or v<0: raise ValueError('Invalid numeric value.')
                facts[key]={'value':v,'status':'reported','source':url,'quote':f'Imported field {key}: {v}'}
            data={'name':name,'url':url,'description':str(row.get('description') or ''),'source':'User-authorized import','facts':facts,'contacts':[],'evidence':[],'export_allowed':str(row.get('export_allowed','false')).lower()=='true','sale_status':'Imported listing; status unverified' if category=='brokered' else 'Seller interest unknown'}
            db.add(Target(user_id=uid(request),job_id=job.id,category=category,identity=identity,data=data,screening=screen(data,criteria.data))); count+=1
        job.progress={category:count,'skipped_rows':len(rows)-count}; db.commit(); return public_job(job)
    except (ValueError,UnicodeError,StopIteration) as e: db.rollback(); raise HTTPException(400,str(e))

class OpportunityInput(BaseModel):
    name:str=Field(min_length=1,max_length=200)
    category:str='proprietary'
    target_id:int|None=None
    criteria_id:int|None=None

@app.post('/api/opportunities')
def create_opportunity(body:OpportunityInput,request:Request,db:DBSession=Depends(session_db)):
    if body.category not in FRAMEWORKS: raise HTTPException(400,'Invalid category.')
    criteria={}; category=body.category; name=body.name
    if body.target_id:
        t=owned(db,Target,body.target_id,uid(request)); criteria=db.get(Job,t.job_id).payload['criteria']; category=t.category; name=t.data['name']
    elif body.criteria_id: criteria=owned(db,CriteriaRecord,body.criteria_id,uid(request)).data
    o=Opportunity(user_id=uid(request),target_id=body.target_id,name=name,category=category,criteria=criteria); db.add(o); db.commit(); return {'id':o.id}

@app.get('/api/opportunities')
def opportunity_list(request:Request,db:DBSession=Depends(session_db)):
    return [{'id':o.id,'name':o.name,'category':o.category} for o in db.scalars(select(Opportunity).where(Opportunity.user_id==uid(request)).order_by(Opportunity.id.desc()))]

@app.get('/api/opportunities/{oid}')
def opportunity_detail(oid:int,request:Request,db:DBSession=Depends(session_db)):
    o=owned(db,Opportunity,oid,uid(request))
    return {'id':o.id,'name':o.name,'category':o.category,'notes':o.notes,'financials':o.financials,'criteria':o.criteria,
            'documents':[{'id':d.id,'filename':d.filename,'size':d.size} for d in db.scalars(select(Document).where(Document.opportunity_id==oid,Document.user_id==uid(request)))],
            'evaluations':[{'id':e.id,'created':e.created,'report':e.report} for e in db.scalars(select(Evaluation).where(Evaluation.opportunity_id==oid,Evaluation.user_id==uid(request)).order_by(Evaluation.id.desc()))]}

def require_idle(db,oid,user_id):
    for j in db.scalars(select(Job).where(Job.user_id==user_id,Job.kind=='evaluation',Job.status.in_(ACTIVE))):
        if j.payload['opportunity_id']==oid: raise HTTPException(409,'An evaluation is active. Wait for completion before changing or deleting its materials.')

class MaterialsInput(BaseModel):
    notes:str=Field('',max_length=40000)
    financials:list[FinancialDatum]=Field(default_factory=list,max_length=100)
    criteria_id:int|None=None

@app.put('/api/opportunities/{oid}')
def update_materials(oid:int,body:MaterialsInput,request:Request,db:DBSession=Depends(session_db)):
    o=owned(db,Opportunity,oid,uid(request)); require_idle(db,oid,uid(request)); o.notes=body.notes
    o.financials=[{**f.model_dump(),'basis':'user_reviewed'} for f in body.financials]
    if body.criteria_id: o.criteria=owned(db,CriteriaRecord,body.criteria_id,uid(request)).data
    db.commit(); return {'ok':True}

@app.post('/api/opportunities/{oid}/documents')
def upload_document(oid:int,request:Request,file:UploadFile=File(...),db:DBSession=Depends(session_db)):
    owned(db,Opportunity,oid,uid(request)); require_idle(db,oid,uid(request))
    count=db.scalar(select(func.count()).select_from(Document).where(Document.opportunity_id==oid))
    if count>=config.MAX_DOCUMENTS: raise HTTPException(400,'Limit is 10 documents. Delete or replace an existing document.')
    raw=file.file.read(config.MAX_UPLOAD+1); filename=Path(file.filename or 'document').name[:200]
    try: key,digest=store(filename,raw)
    except ValueError as e: raise HTTPException(400,str(e))
    d=Document(user_id=uid(request),opportunity_id=oid,filename=filename,storage=key,size=len(raw),digest=digest); db.add(d); db.commit(); return {'id':d.id,'status':'Stored encrypted. No extraction or AI analysis started.'}

@app.post('/api/documents/{did}/preview')
def preview_document(did:int,request:Request,db:DBSession=Depends(session_db)):
    d=owned(db,Document,did,uid(request))
    try: return {'filename':d.filename,'chunks':extract(d.filename,read_document(d))}
    except ValueError as e: raise HTTPException(400,str(e))

@app.delete('/api/documents/{did}')
def delete_document(did:int,request:Request,db:DBSession=Depends(session_db)):
    d=owned(db,Document,did,uid(request)); require_idle(db,d.opportunity_id,uid(request))
    # Prior reports can quote the deleted file; deletion removes those confidential reports too.
    db.execute(delete(Evaluation).where(Evaluation.opportunity_id==d.opportunity_id,Evaluation.user_id==uid(request)))
    (config.DATA/'documents'/d.storage).unlink(missing_ok=True); db.delete(d); db.commit(); return {'ok':True,'prior_reports_removed':True}

class EvalInput(BaseModel):
    idempotency_key:str=Field(min_length=8,max_length=64)
    budget:float=Field(gt=0,le=100,allow_inf_nan=False)
    weights:dict[str,float]|None=None

@app.post('/api/opportunities/{oid}/evaluate')
def evaluate(oid:int,body:EvalInput,request:Request,db:DBSession=Depends(session_db)):
    o=owned(db,Opportunity,oid,uid(request))
    if not config.AI_KEY: raise HTTPException(409,'AI evaluation requires OPENAI_API_KEY in .env. No AI request was made.')
    if body.budget>float(request.state.user.settings.get('evaluation_budget',0.25)): raise HTTPException(400,'Evaluation budget exceeds your configured limit.')
    try: weights=weights_valid(o.category,body.weights or request.state.user.settings.get('frameworks',{}).get(o.category,FRAMEWORKS[o.category]))
    except ValueError as e: raise HTTPException(400,str(e))
    docs=db.scalars(select(Document).where(Document.opportunity_id==oid,Document.user_id==uid(request))).all()
    if not docs and not o.notes and not o.financials: raise HTTPException(400,'Supply documents, notes, or financial information before evaluation.')
    payload={'opportunity_id':oid,'document_ids':[d.id for d in docs],'notes':o.notes,'financials':o.financials,'criteria':o.criteria,'weights':weights}
    j=enqueue(db,request.state.user,'evaluation',payload,body.idempotency_key,body.budget); return public_job(j)

@app.get('/api/evaluations/{eid}/export')
def evaluation_export(eid:int,request:Request,format:str='markdown',db:DBSession=Depends(session_db)):
    e=owned(db,Evaluation,eid,uid(request)); o=db.get(Opportunity,e.opportunity_id)
    if format=='markdown':
        return Response(evaluation_markdown(o.name,e.id,e.report),media_type='text/markdown',headers={'Content-Disposition':f'attachment; filename="investment-evaluation-{e.id}.md"'})
    if format!='json': raise HTTPException(400,'Export format must be markdown or json.')
    return Response(json.dumps({'opportunity':o.name,'version_id':e.id,'created':e.created,'report':e.report},indent=2),media_type='application/json',headers={'Content-Disposition':f'attachment; filename="investment-evaluation-{e.id}.json"'})

@app.delete('/api/opportunities/{oid}')
def delete_opportunity(oid:int,request:Request,db:DBSession=Depends(session_db)):
    o=owned(db,Opportunity,oid,uid(request)); require_idle(db,oid,uid(request))
    for d in db.scalars(select(Document).where(Document.opportunity_id==oid)):
        (config.DATA/'documents'/d.storage).unlink(missing_ok=True); db.delete(d)
    db.execute(delete(Evaluation).where(Evaluation.opportunity_id==oid))
    for j in db.scalars(select(Job).where(Job.user_id==uid(request),Job.kind=='evaluation')):
        if j.payload['opportunity_id']==oid:
            db.execute(delete(Usage).where(Usage.job_id==j.id)); db.delete(j)
    db.delete(o); db.commit(); return {'ok':True}

class SettingsInput(BaseModel):
    monthly_budget:float=Field(100,ge=0,le=10000,allow_inf_nan=False)
    search_budget:float=Field(10,ge=0,le=100,allow_inf_nan=False)
    evaluation_budget:float=Field(.25,ge=0,le=100,allow_inf_nan=False)
    frameworks:dict=Field(default_factory=dict)
    presets:dict=Field(default_factory=dict)

@app.put('/api/settings')
def save_settings(body:SettingsInput,request:Request,db:DBSession=Depends(session_db)):
    try:
        for category,weights in body.frameworks.items():
            if category not in FRAMEWORKS: raise ValueError('Unknown framework category.')
            weights_valid(category,weights)
        if len(body.presets)>20: raise ValueError('At most 20 presets.')
        for name,preset in body.presets.items():
            if len(name)>100 or preset.get('category') not in FRAMEWORKS: raise ValueError('Invalid preset.')
            weights_valid(preset['category'],preset['weights'])
    except (ValueError,TypeError,AttributeError) as e: raise HTTPException(400,str(e))
    user=db.get(User,uid(request)); user.settings=body.model_dump(); db.commit(); return {'ok':True}

@app.get('/api/usage')
def usage(request:Request,db:DBSession=Depends(session_db)):
    return [{'category':u.category,'provider':u.provider,'estimated_cost':u.estimated_cost,'detail':u.detail,'created':u.created,'job_id':u.job_id} for u in db.scalars(select(Usage).where(Usage.user_id==uid(request)).order_by(Usage.id.desc()).limit(500))]

class FactInput(BaseModel):
    key:str=Field(max_length=80)
    value:str|float|bool
    source:str=Field(max_length=2000)
    quote:str=Field(min_length=1,max_length=3000)

@app.post('/api/targets/{tid}/facts')
def verify_fact(tid:int,body:FactInput,request:Request,db:DBSession=Depends(session_db)):
    t=owned(db,Target,tid,uid(request))
    numeric=['revenue','ebitda','sde','asking_price','employees','business_age','ebitda_margin','ebitda_multiple']
    boolean=list(CHARACTERISTICS)+['recurring']
    text=['industry','country','location','currency','customer_type','revenue_model']
    if body.key not in numeric+boolean+text: raise HTTPException(400,'Unknown fact field.')
    if not safe_public_url(body.source): raise HTTPException(400,'Provide a public HTTPS evidence URL.')
    if body.key in boolean and type(body.value) is not bool: raise HTTPException(400,'Use true or false for binary facts.')
    if body.key in numeric and (type(body.value) not in [float,int] or not math.isfinite(body.value)): raise HTTPException(400,'Use a finite numeric value.')
    if body.key in text and (not isinstance(body.value,str) or not body.value.strip()): raise HTTPException(400,'Use a nonempty text value.')
    if body.key=='currency' and body.value not in ['USD','CAD']: raise HTTPException(400,'Supported currencies: USD, CAD.')
    data=dict(t.data); facts=dict(data.get('facts',{})); facts[body.key]={'value':body.value,'status':'verified','source':body.source,'quote':body.quote,'verified_by':'Account holder after evidence review'}; data['facts']=facts; t.data=data
    t.screening=screen(data,db.get(Job,t.job_id).payload['criteria']); db.commit(); return {'screening':t.screening}
