import csv
import io
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

COLUMNS = ['Name / listing title','Category','Description','Source URL','Source','Revenue','EBITDA','SDE','Asking price','Industry','Country','Location','Employees','Screening status','Needs verification','Contact information','Last researched','Sale status','Fit tier','Fit reasons','HQ city','HQ state','Founded year','Employees (Apollo est.)','Revenue (Apollo est., USD)','Financial data basis','Ownership / parent','Funding','Owner / decision-maker name and title','Contact revealed (Y/N)','Apollo URL','LinkedIn URL','Website']

def safe_cell(v):
    value = '' if v is None else str(v)
    return "'"+value if value.lstrip().startswith(('=','+','-','@','\t','\r')) else value


def target_rows(targets):
    rows = []
    tier_order={'Strong fit':0,'Possible fit':1,'Likely out of range':2,'Excluded':3}
    for t in sorted(targets,key=lambda t:(tier_order.get(t.screening.get('tier'),1),str(t.data.get('name','')).casefold())):
        d = t.data
        if not d.get('export_allowed',False): continue
        facts = d.get('facts',{})
        val = lambda k: facts.get(k,{}).get('value','')
        rows.append([d['name'],t.category,d.get('description',''),d.get('url',''),d.get('source',''),val('revenue'),val('ebitda'),val('sde'),val('asking_price'),val('industry'),val('country'),val('location'),val('employees'),t.screening['status'],', '.join(t.screening.get('missing',[])),json.dumps(d.get('contacts',[])),d.get('last_researched',''),d.get('sale_status',''),t.screening.get('tier','Possible fit'),'; '.join(t.screening.get('fit_reasons',[])),val('city'),val('state'),val('founded_year'),val('employees') if d.get('source')=='Apollo' else '',val('revenue') if d.get('source')=='Apollo' else '', '; '.join(f'{key}: {facts[key].get("basis",facts[key].get("status","unknown"))}' for key in ['revenue','ebitda','employees'] if key in facts),val('ownership_parent'),json.dumps({'total_funding':val('total_funding'),'latest_stage':val('latest_funding_stage')}),'; '.join(str(c.get('name') or '')+' / '+str(c.get('title') or '') for c in d.get('contacts',[]) if c.get('name')),'Y' if any(c.get('revealed') for c in d.get('contacts',[])) else 'N',d.get('apollo_url',''),d.get('linkedin_url',''),d.get('website','')])
    return rows


def export_targets(targets,kind):
    rows = target_rows(targets)
    if kind == 'csv':
        buf = io.StringIO(); writer=csv.writer(buf); writer.writerow(COLUMNS)
        writer.writerows([[safe_cell(v) for v in row] for row in rows])
        return buf.getvalue().encode('utf-8-sig')
    wb = Workbook(); ws=wb.active; ws.title='Acquisition targets'; ws.append(COLUMNS)
    for row in rows: ws.append([safe_cell(v) for v in row])
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
    for cell in ws[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='163B40')
    from openpyxl.utils import get_column_letter
    for i in range(1,len(COLUMNS)+1): ws.column_dimensions[get_column_letter(i)].width=25
    buf=io.BytesIO(); wb.save(buf); return buf.getvalue()


def evaluation_markdown(name, version, report):
    lines=[f'# Investment evaluation: {name}',f'Version: {version}',f'Status: {report["status"]}',f'Preliminary score: {report["overall_score"] if report["overall_score"] is not None else "Insufficient Information"}',f'Evidence coverage: {report["evidence_coverage_percent"]}%',f'Confidence: {report["confidence"]}', '', '## Methodology',report['methodology'],'','## Preliminary conclusion',report['conclusion'],'','## Criterion-level assessment']
    for a in report['assessments']:
        lines.extend([f'### {a["criterion"]} ({report["weights"][a["criterion"]]}%)',f'Score: {a["score"] if a["score"] is not None else "Unscored"}',a['explanation'],a['citation_status']])
        for c in a['citations']: lines.append(f'- {c["document"]}, {c["location"]}: {c["quote"]}')
    lines.extend(['','## Financial figures'])
    for f in report['financials']:
        lines.append(f'- {f["metric"]} / {f["year"]}: {f["value"]} {f["currency"]} ({f["basis"]})')
        for c in f['citations']: lines.append(f'  Source: {c["document"]}, {c["location"]}: {c["quote"]}')
    lines.extend(['','## Deterministic calculations',json.dumps(report['calculations'],indent=2)])
    for key in ['strengths','risks','red_flags','missing_information','management_questions','diligence_priorities','extraction_warnings']:
        lines.extend(['',f'## {key.replace("_"," ").title()}']+[f'- {v}' for v in report[key]])
    lines.extend(['','This is preliminary decision support, subject to independent diligence.'])
    return '\n'.join(lines).encode()
